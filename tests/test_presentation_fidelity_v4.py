"""Independent, hash-pinned upstream decks exercise the real presentation routes."""
from __future__ import annotations

import base64
import io
import json
from pathlib import Path

import pytest
from evidence_lane_plugin.document_parsers import digest, package_members
from evidence_lane_plugin.presentation_profile import read_snapshot
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .test_presentation_profile_v4 import call, execute, generation, plan
from .test_presentation_profile_v4 import (
    presentation_system as presentation_system,  # noqa: PLC0414
)

FIXTURES = Path(__file__).parent / 'fixtures/presentations'
CORPUS = [(vendor.name, row['filename']) for vendor in sorted(FIXTURES.iterdir())
          if (vendor / 'manifest.json').is_file()
          for row in json.loads((vendor / 'manifest.json').read_text())['files']
          if Path(row['filename']).suffix.lower() in {'.pptx', '.pptm', '.potx', '.ppsx'}]


def reference(vendor, name):
    manifest = json.loads((FIXTURES / vendor / 'manifest.json').read_text())
    for row in manifest['files']:
        raw = (FIXTURES / vendor / row['filename']).read_bytes()
        assert len(raw) == row['bytes'] and digest(raw) == row['sha256'], row['filename']
    return (FIXTURES / vendor / name).read_bytes()


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def normalize(text):
    return text.replace('\v', '\n')


@pytest.mark.parametrize(('vendor', 'name'), CORPUS)
def test_independent_corpus_intake_and_reader_agree(presentation_system, vendor, name):
    raw = reference(vendor, name)
    store = presentation_system[1]
    path = store.source_root / name
    path.write_bytes(raw)
    plan(presentation_system, ['presentation_index', 'presentation_export'], max_input_bytes=4_194_304)
    indexed = execute(presentation_system, 'presentation_index', {'filename': name})
    manifest, facts = read_snapshot(store, indexed['snapshot_id'])
    assert manifest['source_object'] == digest(raw) == indexed['sha256']
    assert store.lane('ppt').read_object(manifest['raw_object']) == raw
    assert not facts['fidelity']['layout_verified']
    if Path(name).suffix in {'.potx', '.ppsx'}:
        # This reader's factory rejects these content types. Independently read
        # their OPC package order; do not rename/convert bytes to imply support.
        from zipfile import ZipFile

        from lxml import etree
        with ZipFile(io.BytesIO(raw)) as archive:
            root = etree.fromstring(archive.read('ppt/presentation.xml'))
            relationships = etree.fromstring(archive.read('ppt/_rels/presentation.xml.rels'))
        targets = {row.get('Id'): row.get('Target') for row in relationships}
        order = root.xpath('//*[local-name()="sldId"]/@*[local-name()="id" and namespace-uri()!=""]')
        assert len(order) == len(facts['slide_order'])
        assert [targets[identity].split('/')[-1] for identity in order] == [part.split('/')[-1] for part in facts['slide_order']]
    else:
        deck = Presentation(io.BytesIO(raw))
        assert facts['slide_order'] == [str(slide.part.partname).lstrip('/') for slide in deck.slides]
        assert int(facts['slide_size_emu']['cx']) == deck.slide_width
        assert int(facts['slide_size_emu']['cy']) == deck.slide_height
        for number, slide in enumerate(deck.slides, 1):
            items = [item for item in facts['items'] if item.get('slide_number') == number]
            for shape in walk(slide.shapes):
                projected = next(item for item in items if item['kind'] == 'shape'
                                 and not item['is_notes'] and item['shape_id'] == str(shape.shape_id))
                assert projected['name'] == shape.name
                if shape.has_text_frame:
                    assert projected['text'] == normalize(shape.text)
                if projected['geometry']:
                    coords = projected['geometry']['coordinates']
                    assert (int(coords['off']['x']), int(coords['off']['y']),
                            int(coords['ext']['cx']), int(coords['ext']['cy'])) == (
                                shape.left, shape.top, shape.width, shape.height)
                if shape.has_table:
                    expected = [[normalize(cell.text) for cell in row.cells] for row in shape.table.rows]
                    assert any(item['kind'] == 'table' and item['rows'] == expected for item in items)
                    if any(cell.is_merge_origin or cell.is_spanned for cell in shape.table.iter_cells()):
                        assert any(item['kind'] == 'table' and item['merged_cells'] for item in items)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, 'image')):
                    assert any(item['kind'] == 'image' and item['media_sha256'] == digest(shape.image.blob) for item in items)
                if shape.has_chart:
                    cached = [value for item in items if item['kind'] == 'chart' for value in item['cached_values']]
                    numbers = {float(value) for value in cached if value.replace('.', '', 1).replace('-', '', 1).isdigit()}
                    assert {value for series in shape.chart.series for value in series.values if value is not None} <= numbers
            if slide.has_notes_slide:
                text = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ''
                assert any(item['kind'] == 'notes' and normalize(text) in item['text'] for item in items)
    before = {str(p): p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    offset, chunks = 0, []
    while offset is not None:
        read = call(presentation_system, 'presentation_read', {'snapshot_id': indexed['snapshot_id'],
                                                            'offset': offset, 'max_bytes': 131072})
        assert read.status == 'ok', read.error
        returned = read.result['result']
        assert returned['sha256'] == digest(raw) and returned['offset'] == offset
        chunks.append(base64.b64decode(returned['content_base64']))
        offset = returned['next_offset']
    assert b''.join(chunks) == raw
    assert before == {str(p): p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    assert path.read_bytes() == raw
    exported = execute(presentation_system, 'presentation_export', {'snapshot_id': indexed['snapshot_id'],
        'filename': 'exported-' + name}, index=1)
    destination, destination_facts = read_snapshot(store, exported['index_refresh']['result']['snapshot_id'])
    assert not exported['source_index_refresh_required'] and destination_facts == facts
    assert destination['source_path'] == 'exported-' + name and destination['raw_object'] == digest(raw)
    assert (store.source_root / ('exported-' + name)).read_bytes() == raw
    assert not destination_facts['fidelity']['layout_verified']


def test_independent_notes_edit_preserves_all_other_parts(presentation_system):
    raw = reference('python-pptx', 'prs-notes.pptx')
    store = presentation_system[1]
    (store.source_root / 'notes.pptx').write_bytes(raw)
    plan(presentation_system, ['presentation_index', 'presentation_edit'])
    indexed = execute(presentation_system, 'presentation_index', {'filename': 'notes.pptx'})
    _, facts = read_snapshot(store, indexed['snapshot_id'])
    text = next(item for item in facts['items'] if item['kind'] == 'text_block' and item['is_notes'] and item['text'].strip())
    edited = execute(presentation_system, 'presentation_edit', {'snapshot_id': indexed['snapshot_id'],
        'expected_sha256': indexed['sha256'], 'replacements': [{'part': text['part'],
            'paragraph_index': text['ordinal'], 'expected_text': text['text'],
            'replacement_text': 'Independently verified speaker notes.'}]}, index=1)
    changed = store.lane('ppt').read_object(edited['sha256'])
    assert 'Independently verified speaker notes.' in Presentation(io.BytesIO(changed)).slides[0].notes_slide.notes_text_frame.text
    before, after = package_members(raw), package_members(changed)
    assert before.keys() == after.keys()
    assert [name for name in before if before[name] != after[name]] == [text['part']]
    assert (store.source_root / 'notes.pptx').read_bytes() == raw
    historical = call(presentation_system, 'presentation_read', {'snapshot_id': indexed['snapshot_id'], 'max_bytes': len(raw)})
    assert base64.b64decode(historical.result['result']['content_base64']) == raw


def test_independent_theme_styles_and_picture_order_survive_edits(presentation_system):
    raw = reference('python-pptx', 'font-color.pptx')
    store = presentation_system[1]
    (store.source_root / 'colors.pptx').write_bytes(raw)
    picture_raw = reference('python-pptx', 'shp-picture.pptx')
    (store.source_root / 'pictures.pptx').write_bytes(picture_raw)
    plan(presentation_system, ['presentation_index', 'presentation_edit', 'presentation_index', 'presentation_edit'])
    indexed = execute(presentation_system, 'presentation_index', {'filename': 'colors.pptx'})
    edited = execute(presentation_system, 'presentation_edit', {'snapshot_id': indexed['snapshot_id'],
        'expected_sha256': indexed['sha256'], 'replacements': [{'part': 'ppt/slides/slide1.xml',
            'paragraph_index': 2, 'expected_text': 'text with theme color', 'replacement_text': 'Updated theme-colored text'}]}, index=1)
    from lxml import etree
    before = etree.fromstring(package_members(raw)['ppt/slides/slide1.xml'])
    changed = store.lane('ppt').read_object(edited['sha256'])
    after = etree.fromstring(package_members(changed)['ppt/slides/slide1.xml'])
    query = '//*[local-name()="rPr"]'
    assert [etree.tostring(node) for node in before.xpath(query)] == [etree.tostring(node) for node in after.xpath(query)]
    pictures = execute(presentation_system, 'presentation_index', {'filename': 'pictures.pptx'}, index=2)
    reordered = execute(presentation_system, 'presentation_edit', {'snapshot_id': pictures['snapshot_id'],
        'expected_sha256': pictures['sha256'], 'slide_order': ['ppt/slides/slide2.xml', 'ppt/slides/slide1.xml']}, index=3)
    reordered_raw = store.lane('ppt').read_object(reordered['sha256'])
    original_deck, after_deck = Presentation(io.BytesIO(picture_raw)), Presentation(io.BytesIO(reordered_raw))
    assert [shape.image.blob for shape in after_deck.slides[0].shapes] == [shape.image.blob for shape in original_deck.slides[1].shapes]
    assert all(package_members(reordered_raw)[name] == value for name, value in package_members(picture_raw).items()
               if name != 'ppt/presentation.xml')
    assert (store.source_root / 'colors.pptx').read_bytes() == raw
    assert (store.source_root / 'pictures.pptx').read_bytes() == picture_raw


@pytest.mark.parametrize('corruption', ['projection', 'version'])
def test_corrupted_presentation_index_cannot_replace_immutable_truth(presentation_system, corruption):
    plan(presentation_system, ['presentation_index'])
    indexed = execute(presentation_system, 'presentation_index', {'filename': 'fixture.pptx'})
    lane = presentation_system[1].lane('ppt')
    with lane.transaction() as db:
        if corruption == 'projection':
            db.execute("UPDATE ppt_shape SET payload_json=json_set(payload_json,'$.name','forged')")
        else:
            db.execute('UPDATE ppt_version SET generation=99 WHERE snapshot_id=?', (indexed['snapshot_id'],))
    response = call(presentation_system, 'presentation_query', {'snapshot_id': indexed['snapshot_id'], 'collection': 'shape'})
    assert response.status == 'error'
    assert response.error.code == ('PRESENTATION_QUERY_INTEGRITY' if corruption == 'projection' else 'PRESENTATION_SNAPSHOT_INTEGRITY')


def test_generated_image_shapes_geometry_and_multiline_text_are_editable(presentation_system):
    raw = reference('python-pptx', 'shp-picture.pptx')
    image = Presentation(io.BytesIO(raw)).slides[0].shapes[0].image
    request = generation(1)
    request['slides'][0]['objects'] += [
        {'kind': 'image', 'name': 'Exact embedded image', 'x': 11, 'y': 3, 'width': 1, 'height': 1,
         'image_format': 'jpeg' if image.ext == 'jpg' else image.ext,
         'image_base64': base64.b64encode(image.blob).decode(), 'description': 'Preserved upstream fixture image'},
        {'kind': 'rectangle', 'name': 'Rectangle', 'x': .5, 'y': 6, 'width': 2, 'height': .8,
         'text': 'Two\nlines', 'fill': 'DDEEFF'},
        {'kind': 'ellipse', 'name': 'Ellipse', 'x': 3, 'y': 6, 'width': 2, 'height': .8, 'text': 'Editable'},
    ]
    plan(presentation_system, ['presentation_generate'])
    generated = execute(presentation_system, 'presentation_generate', request)
    deck = Presentation(io.BytesIO(presentation_system[1].lane('ppt').read_object(generated['sha256'])))
    shapes = deck.slides[0].shapes
    assert shapes[3].image.blob == image.blob
    assert (shapes[3].left, shapes[3].top, shapes[3].width, shapes[3].height) == tuple(round(value * 914400) for value in (11, 3, 1, 1))
    assert shapes[4].text == 'Two\nlines' and shapes[5].text == 'Editable'
    assert str(shapes[4].auto_shape_type).startswith('RECTANGLE')
    assert str(shapes[5].auto_shape_type).startswith('OVAL')
