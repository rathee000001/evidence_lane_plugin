"""Actual SDK, Delta, OS-worker and separate-lane tabular integration."""
from __future__ import annotations

import base64
import hashlib
import io
import json
import time
import zipfile

import pytest
from evidence_lane_plugin.connections import ConnectRequest, ProjectSelection
from evidence_lane_plugin.engine import Engine
from evidence_lane_plugin.errors import LaneError
from evidence_lane_plugin.internal_sdk import PublicActionSDKDispatcher
from evidence_lane_plugin.plan_runtime import PlanCreate, PlanStore, TaskBudget, TaskDefinition
from evidence_lane_plugin.sdk import ActionRequest
from evidence_lane_plugin.tabular_workers import tabular_worker_operations
from evidence_lane_plugin.workers import WorkerOperation, WorkerPool


def spreadsheet_bytes():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table, TableStyleInfo
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Inputs'
    sheet.append(['Item', 'Quantity', 'Price', 'Total'])
    sheet.append(['Alpha', 3, 2.5, '=B2*C2'])
    sheet.append(['Beta', 2, 4, '=B3*C3'])
    sheet['A2'].font = Font(bold=True, color='123456')
    table = Table(displayName='Items', ref='A1:D3')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    summary = workbook.create_sheet('Summary')
    summary['A1'], summary['B1'] = 'Total', "=SUM('Inputs'!D2:D3)"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.title = 'Quantities'
    summary.add_chart(chart, 'A3')
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.fixture
def tabular_system(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'fixture.xlsx').write_bytes(spreadsheet_bytes())
    (source / 'values.json').write_text('[{"id":"0007","value":0.123456789012345678901,"active":true},'
        '{"id":"0008","value":2.75,"active":false},{"id":"0009","value":null,"active":true}]', encoding='utf-8')
    operations = (*tabular_worker_operations(), WorkerOperation('render_lane_view',
        'evidence_lane_plugin.artifact_contract', 'render_lane_view_worker', dependencies=('langgraph', 'langchain_core', 'graphviz')))
    with Engine(tmp_path / 'runtime', worker_pool=WorkerPool(operations, workers=1)) as engine:
        entry = engine.directory.register(tmp_path / 'state', source_root=source, create=True, read_only=False)
        store = engine.directory.open(entry['project_id'], write=True)
        _, session = engine.clients.connect(ConnectRequest(projects=[ProjectSelection(project_id=store.project_id,
            permissions=['read', 'write', 'tools', 'admin'])]))
        yield engine, store, session


def call(system, action, arguments=None, **kwargs):
    engine, store, session = system
    return PublicActionSDKDispatcher(engine).execute(ActionRequest(action=action, project_id=store.project_id,
        arguments=arguments or {}, **kwargs), session)


def plan(system, actions, *, max_input_bytes=1_048_576):
    engine, store, _ = system
    tasks = [TaskDefinition(task_id='table-' + str(index), title=action, requested_outcome='Verify the selected file operation',
        profile=engine.registry.get(action).profile, allowed_actions=[action], permitted_tools=['Python', 'openpyxl', 'lxml',
            'python_calamine', 'pyarrow', 'pandas', 'Polars', 'DuckDB', 'LibreOffice', 'pypdfium2',
            'LangGraph_Mermaid_engine', 'Python_Graphviz_DOT_engine', 'rustworkx'], permitted_paths=['.'],
        acceptance_checks=list(engine.registry.get(action).verification_checks),
        budget=TaskBudget(max_input_bytes=max_input_bytes, max_output_bytes=33_554_432, max_seconds=180)) for index, action in enumerate(actions)]
    with engine.project_work.mutation(store) as lease:
        PlanStore(store).create(PlanCreate(title='Tabular fixture', tasks=tasks), lease, actor_id='fixture')


def execute(system, action, arguments, *, index=0, expected='verified'):
    store = system[1]
    task = PlanStore(store).task('table-' + str(index), expected_revision=1)
    admitted = call(system, 'delta_enter', {'task_id': task.definition.task_id, 'plan_revision': 1,
        'contract_digest': task.contract_digest, 'action': action, 'arguments': arguments}, expected_revision=1)
    assert admitted.status == 'queued', admitted.error
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        try:
            with store.lane('plan').connection(read_only=True) as connection:
                row = dict(connection.execute('SELECT * FROM delta_runs WHERE job_id=?', (admitted.job_id,)).fetchone())
        except LaneError as error:
            if error.code != 'PROJECT_RECOVERY_REQUIRED':
                raise
            time.sleep(.02)
            continue
        if row['state'] in {'verified', 'blocked'}:
            break
        time.sleep(.02)
    assert row['state'] == expected, {key: row[key] for key in ('state', 'error_code', 'result_object')}
    system[0].delta.owned_completion(admitted.job_id).result(
        timeout=max(1, deadline - time.monotonic())
    )
    if expected == 'blocked':
        return row
    return json.loads(store.lane('plan').read_object(row['result_object']))['result']['result']


def query(system, prefix, snapshot, **arguments):
    response = call(system, prefix + '_query', {'snapshot_id': snapshot, **arguments})
    assert response.status == 'ok', response.error
    return response.result['result']


def test_native_workbook_facts_and_data_are_in_separate_databases(tabular_system):
    plan(tabular_system, ['spreadsheet_index', 'data_index'])
    workbook = execute(tabular_system, 'spreadsheet_index', {'filename': 'fixture.xlsx'})
    data = execute(tabular_system, 'data_index', {'filename': 'values.json'}, index=1)
    store = tabular_system[1]
    assert '/data_excel/' in workbook['natural_path'].replace('\\', '/')
    assert '/data/' in data['natural_path'].replace('\\', '/')
    formulas = query(tabular_system, 'spreadsheet', workbook['snapshot_id'], collection='formula')['rows']
    summary = next(row for row in formulas if row['sheet'] == 'Summary')
    assert summary['formula'] == "SUM('Inputs'!D2:D3)"
    assert summary['dependencies'] == [{'sheet': 'Inputs', 'range': 'D2:D3'}]
    assert summary['cached_value']['type'] == 'null'
    assert len(query(tabular_system, 'spreadsheet', workbook['snapshot_id'], collection='table')['rows']) == 1
    assert len(query(tabular_system, 'spreadsheet', workbook['snapshot_id'], collection='chart')['rows']) == 1
    assert query(tabular_system, 'spreadsheet', workbook['snapshot_id'], query='Alpha')['rows'][0]['text'] == 'Alpha'
    records = query(tabular_system, 'data', data['snapshot_id'], collection='row')['rows']
    assert records[0]['values'][1] == {'type': 'text', 'value': '0007'}
    assert records[0]['values'][2] == {'type': 'number', 'value': '0.123456789012345678901'}
    with store.lane('data_excel').connection(read_only=True) as connection:
        assert not connection.execute("SELECT 1 FROM sqlite_schema WHERE name='data_table'").fetchone()
    with store.lane('data').connection(read_only=True) as connection:
        assert not connection.execute("SELECT 1 FROM sqlite_schema WHERE name='sheet_workbook'").fetchone()
    wrong_lane = call(tabular_system, 'data_read', {'snapshot_id': workbook['snapshot_id']})
    assert wrong_lane.status == 'error'
    assert wrong_lane.error.code == 'TABULAR_SNAPSHOT_MISSING'


def test_cell_edit_preserves_other_parts_and_exact_export(tabular_system):
    plan(tabular_system, ['spreadsheet_index', 'spreadsheet_edit', 'spreadsheet_export'])
    store = tabular_system[1]
    original = (store.source_root / 'fixture.xlsx').read_bytes()
    indexed = execute(tabular_system, 'spreadsheet_index', {'filename': 'fixture.xlsx'})
    edited = execute(tabular_system, 'spreadsheet_edit', {'snapshot_id': indexed['snapshot_id'], 'expected_sha256': indexed['sha256'],
        'replacements': [{'sheet': 'Inputs', 'cell': 'B2', 'expected_value': {'type': 'number', 'value': '3'},
                          'replacement_value': {'type': 'number', 'value': '5'}}]}, index=1)
    assert edited['tool_evidence']['untouched_members_byte_identical']
    cell = next(row for row in query(tabular_system, 'spreadsheet', edited['snapshot_id'], collection='cell')['rows'] if row['sheet'] == 'Inputs' and row['cell'] == 'B2')
    assert cell['value'] == {'type': 'number', 'value': '5'}
    assert (store.source_root / 'fixture.xlsx').read_bytes() == original
    exported = execute(tabular_system, 'spreadsheet_export', {'snapshot_id': edited['snapshot_id'], 'filename': 'output.xlsx'}, index=2)
    output = (store.source_root / 'output.xlsx').read_bytes()
    assert hashlib.sha256(output).hexdigest() == exported['after_sha256']
    with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
        for name in before.namelist():
            if name not in edited['tool_evidence']['changed_parts']:
                assert before.read(name) == after.read(name)
    historical = call(tabular_system, 'spreadsheet_read', {'snapshot_id': indexed['snapshot_id'], 'max_bytes': 131072})
    assert base64.b64decode(historical.result['result']['content_base64']) == original


def test_data_transform_exact_decimal_types_and_lineage(tabular_system):
    plan(tabular_system, ['data_index', 'data_transform'])
    source = execute(tabular_system, 'data_index', {'filename': 'values.json'})
    transformed = execute(tabular_system, 'data_transform', {'snapshot_id': source['snapshot_id'], 'expected_sha256': source['sha256'],
        'logical_name': 'active.json', 'filters': [{'column': 'active', 'operator': 'equals', 'value': {'type': 'boolean', 'value': True}}],
        'columns': ['id', 'value']}, index=1)
    rows = query(tabular_system, 'data', transformed['snapshot_id'], collection='row')['rows']
    assert len(rows) == 2
    assert rows[0]['values'][1] == {'type': 'number', 'value': '0.123456789012345678901'}
    assert rows[1]['values'][1] == {'type': 'null', 'value': None}
    metadata = query(tabular_system, 'data', transformed['snapshot_id'], collection='metadata')
    assert metadata['inputs'][0]['snapshot_id'] == source['snapshot_id']
    lineage = query(tabular_system, 'data', transformed['snapshot_id'], collection='lineage')['rows']
    assert lineage[0]['input']['snapshot_id'] == source['snapshot_id']


def test_generate_formula_and_formula_like_text_are_distinct(tabular_system):
    plan(tabular_system, ['spreadsheet_generate'])
    generated = execute(tabular_system, 'spreadsheet_generate', {'logical_name': 'created.xlsx', 'sheets': [{'name': 'Results', 'cells': [
        {'cell': 'A1', 'value': {'type': 'text', 'value': 'Quantity'}},
        {'cell': 'A2', 'value': {'type': 'number', 'value': '7'}},
        {'cell': 'B2', 'value': {'type': 'null', 'value': None}, 'formula': 'A2*2'},
        {'cell': 'C2', 'value': {'type': 'text', 'value': '=A2*3'}}]}]})
    rows = query(tabular_system, 'spreadsheet', generated['snapshot_id'], collection='cell')['rows']
    cells = {row['cell']: row for row in rows}
    assert cells['B2']['formula'] == 'A2*2'
    assert cells['B2']['cached_value']['type'] == 'null'
    assert cells['C2']['formula'] is None
    assert cells['C2']['value'] == {'type': 'text', 'value': '=A2*3'}


def test_public_schemas_fix_owning_lane(tabular_system):
    wrong_lane = call(tabular_system, 'spreadsheet_current', {'lane_id': 'data'})
    assert wrong_lane.status == 'error'
    assert call(tabular_system, 'spreadsheet_current').status == 'ok'
