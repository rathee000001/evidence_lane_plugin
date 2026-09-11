"""Artifact production-to-intake over real stdio and the canonical worker pool."""
import asyncio
import base64
import hashlib
import io
import json
import os
from pathlib import Path

from evidence_lane_plugin.engine_runtime import create_runtime_engine
from evidence_lane_plugin.local_transport import LocalEndpoint
from evidence_lane_plugin.plan_runtime import PlanCreate, PlanStore, TaskBudget, TaskDefinition
from evidence_lane_plugin.runtime_health import CapabilityMonitor
from evidence_lane_plugin.sector_evidence_profile import read_snapshot

from .test_native_workflow_bindings import native
from .test_presentation_profile_v4 import generation


def test_stdio_generated_exported_artifact_refresh_and_provenance(tmp_path):
    from pptx import Presentation

    source = tmp_path / 'source'
    source.mkdir()
    with create_runtime_engine(tmp_path / 'runtime', workers=1, capabilities=CapabilityMonitor()) as engine:
        record = engine.directory.register(tmp_path / 'state', source_root=source, create=True, read_only=False)
        store = engine.directory.open(record['project_id'], write=True)
        actions = ['presentation_generate', 'presentation_export', 'artifacts_index',
            'presentation_generate', 'presentation_export', 'artifacts_refresh']
        tasks = [TaskDefinition(task_id='artifact-' + str(index), title=action,
            requested_outcome='Keep exact produced artifact bytes and attributed history',
            profile='artifacts' if action.startswith('artifacts_') else 'presentation',
            allowed_actions=[action], permitted_paths=['.'], permitted_tools=[
                'Python', 'PPTX_OpenXML', 'Pillow', 'lxml', 'SQLite_FTS5_BM25',
                'LangGraph_Mermaid_engine', 'Python_Graphviz_DOT_engine', 'rustworkx'],
            acceptance_checks=list(engine.registry.get(action).verification_checks),
            budget=TaskBudget(max_seconds=120, max_output_bytes=33_554_432))
            for index, action in enumerate(actions)]
        with engine.project_work.mutation(store) as lease:
            PlanStore(store).create(PlanCreate(title='Artifact production fixture', tasks=tasks), lease, actor_id='fixture')

        def join():
            with engine._admission:
                assert engine._admission.wait_for(lambda: engine._background_jobs == 0, timeout=120)

        async def run():
            async with native(engine.root, store.project_id, permissions=('read', 'write', 'tools')) as session:
                async def call(action, arguments=None):
                    response = (await session.call_tool(action, {'project_id': store.project_id,
                        'arguments': arguments or {}})).structuredContent
                    assert response['status'] == 'ok', json.dumps(response, indent=2)
                    return response['result']

                async def execute(index, arguments):
                    task = PlanStore(store).task('artifact-' + str(index), expected_revision=1)
                    admitted = (await session.call_tool('delta_enter', {'project_id': store.project_id,
                        'expected_revision': 1, 'arguments': {'task_id': task.definition.task_id,
                            'plan_revision': 1, 'contract_digest': task.contract_digest,
                            'action': actions[index], 'arguments': arguments}})).structuredContent
                    assert admitted['status'] == 'queued', admitted
                    await asyncio.to_thread(join)
                    with store.lane('plan').connection(read_only=True) as connection:
                        row = dict(connection.execute('SELECT * FROM delta_runs WHERE job_id=?', (admitted['job_id'],)).fetchone())
                    assert row['state'] == 'verified', row
                    value = json.loads(store.lane('plan').read_object(row['result_object']))['result']['result']
                    return value, admitted['job_id'], task.definition.task_id

                assert (await call('client_context'))['native_task_attestation'] == 'not_provided'
                artifacts, producers, previous_export, prior_view = [], [], None, None
                for version, word in enumerate(['alpha', 'beta']):
                    body = generation(slides=1)
                    body['slides'][0]['objects'][0]['text'] = 'Produced artifact ' + word
                    body['expected_snapshot'] = producers[-1]['snapshot_id'] if producers else None
                    produced, _, _ = await execute(version * 3, body)
                    raw = store.lane('ppt').read_object(produced['sha256'])
                    independent = Presentation(io.BytesIO(raw))
                    assert independent.slides[0].shapes[0].text == 'Produced artifact ' + word
                    assert Path(produced['natural_path']).read_bytes() == raw
                    assert not produced['fidelity']['layout_verified']
                    exported, _, _ = await execute(version * 3 + 1, {'snapshot_id': produced['snapshot_id'],
                        'filename': 'delivered.pptx', 'expected_sha256': previous_export})
                    assert exported['after_sha256'] == hashlib.sha256(raw).hexdigest() == produced['sha256']
                    assert (source / 'delivered.pptx').read_bytes() == raw
                    arguments = {'filename': 'delivered.pptx',
                        'expected_snapshot': artifacts[-1]['snapshot_id'] if artifacts else None}
                    indexed, job_id, task_id = await execute(version * 3 + 2, arguments)
                    manifest, facts = read_snapshot(store, 'artifacts', indexed['snapshot_id'])
                    assert manifest['raw_object'] == produced['sha256']
                    assert manifest['task_id'] == task_id and manifest['job_id'] == job_id and manifest['plan_revision'] == 1
                    assert manifest['worker_envelope']['worker_pid'] != os.getpid()
                    assert not facts['fidelity']['source_assertions_validated']
                    assert any(row['kind'] == 'artifact_text_extract' and row.get('native_item_id') for row in facts['items'])
                    assert manifest['previous_snapshot'] == (artifacts[-1]['snapshot_id'] if artifacts else None)
                    if prior_view:
                        state = await call('lane_view_read', {'view_id': 'artifacts.structure'})
                        assert state['state'] == 'stale'
                    selected = await call('lane_view_preview', {'view_id': 'artifacts.structure',
                        'scope': {'query': indexed['source_id']}})
                    view = await call('lane_view_refresh', {'view_id': 'artifacts.structure',
                        'scope': selected['scope'], 'contract_digest': selected['contract_digest'],
                        'source_digest': selected['source_digest'], 'expected_generation': selected['generation'],
                        'formats': ['mmd', 'dot'], 'include_pointer': True})
                    assert {Path(row['path']).name for row in view['files']} == {'artifacts.mmd', 'artifacts.dot', 'lane_pointer.json'}
                    assert all('/sectors/artifacts/' in row['path'].replace('\\', '/') for row in view['files'])
                    artifacts.append(indexed | {'bytes': raw, 'word': word, 'view': view})
                    producers.append(produced)
                    previous_export, prior_view = exported['after_sha256'], view

                before = {str(path): path.read_bytes() for path in store.root.rglob('*.sqlite')}
                for artifact in artifacts:
                    fetched = (await call('lane_fetch', {'lane_id': 'artifacts', 'snapshot_id': artifact['snapshot_id'],
                        'path': 'delivered.pptx', 'representation': 'original_source', 'max_bytes': 131072}))['result']['read']['result']
                    assert base64.b64decode(fetched['content_base64']) == artifact['bytes']
                    searched = (await call('artifacts_query', {'snapshot_id': artifact['snapshot_id'],
                        'query': artifact['word']}))['result']
                    assert searched['rows']
                    historical = await call('lane_view_read', {'view_id': 'artifacts.structure',
                        'snapshot_digest': artifact['view']['snapshot_digest'], 'include_content': True})
                    pointer = json.loads(historical['contents']['pointer'])
                    assert {row['snapshot_id'] for row in pointer['artifacts_and_parts'].values()} == {artifact['snapshot_id']}
                assert (await call('artifacts_current'))['result']['files'][0]['snapshot_id'] == artifacts[-1]['snapshot_id']
                assert before == {str(path): path.read_bytes() for path in store.root.rglob('*.sqlite')}
                assert not (store.root / 'sectors/research').exists() and not (store.root / 'sectors/custom').exists()
                with store.connection(read_only=True) as connection:
                    assert connection.execute("SELECT count(*) FROM sqlite_schema WHERE name LIKE 'artifact_%' OR name LIKE 'ppt_%'").fetchone()[0] == 0
        with LocalEndpoint(engine, studio_enabled=False):
            asyncio.run(run())
