"""Document profile qualification through the real Delta and OS-worker paths."""
from __future__ import annotations

import base64
import hashlib
import io
import json
import time
import zipfile

import pytest
from evidence_lane_plugin.connections import ConnectRequest, ProjectSelection
from evidence_lane_plugin.document_workers import document_worker_operations
from evidence_lane_plugin.engine import Engine
from evidence_lane_plugin.errors import LaneError
from evidence_lane_plugin.internal_sdk import PublicActionSDKDispatcher
from evidence_lane_plugin.plan_runtime import PlanCreate, PlanStore, TaskBudget, TaskDefinition
from evidence_lane_plugin.sdk import ActionRequest
from evidence_lane_plugin.workers import WorkerOperation, WorkerPool


@pytest.fixture
def document_system(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    from docx import Document
    document = Document()
    document.add_heading('Document fixture', 0)
    document.add_heading('Scope', 1)
    document.add_paragraph('The selected operation preserves the source document.')
    document.add_paragraph('A second paragraph has ')
    document.paragraphs[-1].add_run('bold formatting.').bold = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Property', 'Value'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Owner', 'Docs lane'
    document.sections[0].header.paragraphs[0].text = 'Fixture header'
    document.sections[0].footer.paragraphs[0].text = 'Fixture footer'
    document.save(source / 'fixture.docx')
    operations = (*document_worker_operations(), WorkerOperation('render_lane_view',
        'evidence_lane_plugin.artifact_contract', 'render_lane_view_worker', dependencies=('langgraph', 'langchain_core', 'graphviz')))
    with Engine(tmp_path / 'runtime', worker_pool=WorkerPool(operations, workers=1)) as engine:
        entry = engine.directory.register(tmp_path / 'state', source_root=source, create=True, read_only=False)
        store = engine.directory.open(entry['project_id'], write=True)
        _, session = engine.clients.connect(ConnectRequest(projects=[ProjectSelection(project_id=store.project_id,
            permissions=['read', 'write', 'tools', 'admin'])]))
        yield engine, store, session


def call(system, action, arguments=None, **kwargs):
    engine, store, session = system
    return PublicActionSDKDispatcher(engine).execute(ActionRequest(action=action, project_id=store.project_id,
        arguments=arguments or {}, **kwargs), session)


def create_plan(system, actions, *, max_input_bytes=1_048_576, **task_options):
    engine, store, _ = system
    tasks = [TaskDefinition(task_id='doc-' + str(i), title=action, requested_outcome='Verify the selected document operation',
        profile='document', allowed_actions=[action], permitted_tools=['Python', 'SQLite_FTS5_BM25',
            'DOCX_OpenXML', 'lxml', 'LibreOffice', 'pypdfium2', 'Docling',
            'LangGraph_Mermaid_engine', 'Python_Graphviz_DOT_engine', 'rustworkx'], permitted_paths=['.'],
        acceptance_checks=list(engine.registry.get(action).verification_checks),
        budget=TaskBudget(max_input_bytes=max_input_bytes, max_output_bytes=33_554_432, max_seconds=180), **task_options)
        for i, action in enumerate(actions)]
    with engine.project_work.mutation(store) as lease:
        PlanStore(store).create(PlanCreate(title='Document fixture', tasks=tasks), lease, actor_id='fixture')


def execute(system, action, arguments, *, index=0, timeout=60):
    store = system[1]
    task = PlanStore(store).task('doc-' + str(index), expected_revision=1)
    admitted = call(system, 'delta_enter', {'task_id': task.definition.task_id, 'plan_revision': 1,
        'contract_digest': task.contract_digest, 'action': action, 'arguments': arguments}, expected_revision=1)
    assert admitted.status == 'queued', admitted.error
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            with store.lane('plan').connection(read_only=True) as connection:
                row = dict(connection.execute('SELECT * FROM delta_runs WHERE job_id=?', (admitted.job_id,)).fetchone())
        except LaneError as error:
            if error.code != 'PROJECT_RECOVERY_REQUIRED':
                raise
            time.sleep(.02)
            continue
        if row['state'] in {'verified', 'blocked'}:
            break
        time.sleep(.02)
    else:
        pytest.fail('Document operation did not reach a readable terminal state within its deadline')
    assert row['state'] == 'verified', {key: row[key] for key in ('state', 'error_code', 'result_object')}
    return json.loads(store.lane('plan').read_object(row['result_object']))['result']['result']


def test_native_intake_queries_original_bytes_and_separate_docs_store(document_system):
    create_plan(document_system, ['document_index'])
    store = document_system[1]
    original = (store.source_root / 'fixture.docx').read_bytes()
    indexed = execute(document_system, 'document_index', {'filename': 'fixture.docx'})
    assert indexed['sha256'] == hashlib.sha256(original).hexdigest()
    assert '/sectors/docs/files/natural/' in indexed['natural_path'].replace('\\', '/')
    assert not indexed['source_bytes_mutated'] and indexed['fidelity']['layout'] == 'not_rendered'
    args = {'snapshot_id': indexed['snapshot_id']}
    before = {p: p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    text = call(document_system, 'document_query', args | {'query': 'preserves source'})
    assert text.status == 'ok', text.error
    assert text.result['result']['rows'][0]['text'] == 'The selected operation preserves the source document.'
    table = call(document_system, 'document_query', args | {'collection': 'table'})
    assert table.result['result']['rows'][0]['rows'][1] == ['Owner', 'Docs lane']
    heading = call(document_system, 'document_query', args | {'collection': 'heading'})
    assert heading.result['result']['rows'][0]['text'] == 'Scope'
    raw = call(document_system, 'document_read', args | {'max_bytes': 131072})
    assert base64.b64decode(raw.result['result']['content_base64']) == original
    after = {p: p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    assert before == after
    assert (store.source_root / 'fixture.docx').read_bytes() == original
    with store.connection(read_only=True) as db:
        assert db.execute("SELECT count(*) FROM sqlite_schema WHERE name GLOB 'doc_*'").fetchone()[0] == 0
    with store.lane('docs').connection(read_only=True) as db:
        assert db.execute('SELECT count(*) FROM doc_table_extract').fetchone()[0] == 1


def test_tracked_edit_preserves_other_package_members_and_original_until_export(document_system):
    create_plan(document_system, ['document_index', 'document_edit', 'document_export', 'document_refresh'])
    store = document_system[1]
    original = (store.source_root / 'fixture.docx').read_bytes()
    indexed = execute(document_system, 'document_index', {'filename': 'fixture.docx'})
    edited = execute(document_system, 'document_edit', {'snapshot_id': indexed['snapshot_id'], 'expected_sha256': indexed['sha256'],
        'replacements': [{'paragraph_index': 2, 'expected_text': 'The selected operation preserves the source document.',
            'replacement_text': 'The edited version is a separate document.', 'tracked': True}]}, index=1)
    assert edited['generation'] == 2 and edited['tool_evidence']['untouched_members_byte_identical']
    assert (store.source_root / 'fixture.docx').read_bytes() == original
    rows = call(document_system, 'document_query', {'snapshot_id': edited['snapshot_id'], 'collection': 'deletion'})
    assert rows.result['result']['rows'][0]['text'] == 'The selected operation preserves the source document.'
    edited_bytes = store.lane('docs').read_object(edited['sha256'])
    with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(edited_bytes)) as after:
        for name in before.namelist():
            if name != 'word/document.xml':
                assert before.read(name) == after.read(name)
    exported = execute(document_system, 'document_export', {'snapshot_id': edited['snapshot_id'], 'filename': 'fixture.docx',
        'expected_sha256': indexed['sha256']}, index=2)
    assert exported['source_bytes_mutated'] and (store.source_root / 'fixture.docx').read_bytes() == edited_bytes
    destination = exported['index_refresh']['result']
    assert not exported['source_index_refresh_required'] and destination['generation'] == 3
    refreshed = execute(document_system, 'document_refresh', {'filename': 'fixture.docx', 'expected_snapshot': destination['snapshot_id']}, index=3)
    assert refreshed['generation'] == 4 and refreshed['sha256'] == edited['sha256']
    historical = call(document_system, 'document_read', {'snapshot_id': indexed['snapshot_id'], 'max_bytes': 131072})
    assert base64.b64decode(historical.result['result']['content_base64']) == original


def test_generation_and_new_file_export(document_system):
    create_plan(document_system, ['document_generate', 'document_export'])
    generated = execute(document_system, 'document_generate', {'logical_name': 'report.docx', 'title': 'Document operations',
        'blocks': [{'kind': 'heading', 'text': 'Scope'}, {'kind': 'paragraph', 'text': 'A generated report with a bounded table.'},
                   {'kind': 'table', 'rows': [['Operation', 'Result'], ['Generate', 'Versioned document']]}]})
    assert generated['fidelity']['layout'] == 'not_rendered'
    exported = execute(document_system, 'document_export', {'snapshot_id': generated['snapshot_id'], 'filename': 'new-report.docx'}, index=1)
    assert exported['before_sha256'] is None
    assert hashlib.sha256((document_system[1].source_root / 'new-report.docx').read_bytes()).hexdigest() == generated['sha256']
