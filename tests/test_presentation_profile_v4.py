"""PPT operations through the registered Delta, lane store and OS-worker paths."""
from __future__ import annotations

import base64
import io
import json
import time

import pytest
from evidence_lane_plugin.connections import ConnectRequest, ProjectSelection
from evidence_lane_plugin.document_parsers import digest, package_members
from evidence_lane_plugin.engine import Engine
from evidence_lane_plugin.errors import LaneError
from evidence_lane_plugin.internal_sdk import PublicActionSDKDispatcher
from evidence_lane_plugin.plan_runtime import PlanCreate, PlanStore, TaskBudget, TaskDefinition
from evidence_lane_plugin.presentation_authoring import generate_package
from evidence_lane_plugin.presentation_workers import presentation_worker_operations
from evidence_lane_plugin.sdk import ActionRequest
from evidence_lane_plugin.workers import WorkerOperation, WorkerPool


def generation(slides=2):
    return {'logical_name': 'fixture.pptx', 'title': 'Presentation operations', 'slides': [
        {'name': f'Slide {index}', 'notes': f'Speaker evidence for slide {index}.', 'objects': [
            {'kind': 'text', 'name': 'Title', 'x': .5, 'y': .5, 'width': 12, 'height': .8,
             'text': f'Versioned presentation {index}', 'font_points': 32, 'bold': True},
            {'kind': 'text', 'name': 'Body', 'x': .5, 'y': 1.5, 'width': 12, 'height': 1,
             'text': 'The selected operation preserves the original source.'},
            {'kind': 'table', 'name': 'Evidence table', 'x': .5, 'y': 3, 'width': 10, 'height': 2,
             'font_points': 18, 'rows': [['Operation', 'Result'], ['Intake', 'Exact source bytes'], ['Edit', 'New immutable version']]},
        ]} for index in range(1, slides + 1)]}


@pytest.fixture
def presentation_system(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'fixture.pptx').write_bytes(generate_package(generation()))
    operations = (*presentation_worker_operations(), WorkerOperation('render_lane_view',
        'evidence_lane_plugin.artifact_contract', 'render_lane_view_worker', dependencies=('langgraph', 'langchain_core', 'graphviz')))
    with Engine(tmp_path / 'runtime', worker_pool=WorkerPool(operations, workers=1)) as engine:
        entry = engine.directory.register(tmp_path / 'state', source_root=source, create=True, read_only=False)
        store = engine.directory.open(entry['project_id'], write=True)
        _, session = engine.clients.connect(ConnectRequest(projects=[ProjectSelection(project_id=store.project_id,
            permissions=['read', 'write', 'tools', 'admin'])]))
        yield engine, store, session


def call(system, action, arguments=None, **kwargs):
    engine, store, session = system
    return PublicActionSDKDispatcher(engine).execute(ActionRequest(action=action, project_id=store.project_id,
        arguments=arguments or {}, **kwargs), session)


def plan(system, actions, *, max_input_bytes=1_048_576):
    engine, store, _ = system
    tasks = [TaskDefinition(task_id='ppt-' + str(index), title=action, requested_outcome='Verify the exact presentation operation',
        profile='presentation', allowed_actions=[action], permitted_tools=['Python', 'SQLite_FTS5_BM25',
            'PPTX_OpenXML', 'Pillow', 'lxml', 'LibreOffice', 'pypdfium2', 'Docling',
            'LangGraph_Mermaid_engine', 'Python_Graphviz_DOT_engine', 'rustworkx'], permitted_paths=['.'],
        acceptance_checks=list(engine.registry.get(action).verification_checks),
        budget=TaskBudget(max_input_bytes=max_input_bytes, max_output_bytes=33_554_432, max_seconds=180)) for index, action in enumerate(actions)]
    with engine.project_work.mutation(store) as lease:
        PlanStore(store).create(PlanCreate(title='Presentation fixture', tasks=tasks), lease, actor_id='fixture')


def execute(system, action, arguments, *, index=0, timeout=60):
    store = system[1]
    task = PlanStore(store).task('ppt-' + str(index), expected_revision=1)
    admitted = call(system, 'delta_enter', {'task_id': task.definition.task_id, 'plan_revision': 1,
        'contract_digest': task.contract_digest, 'action': action, 'arguments': arguments}, expected_revision=1)
    assert admitted.status == 'queued', admitted.error
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            with store.lane('plan').connection(read_only=True) as connection:
                row = dict(connection.execute('SELECT * FROM delta_runs WHERE job_id=?', (admitted.job_id,)).fetchone())
        except LaneError as error:
            if error.code != 'PROJECT_RECOVERY_REQUIRED':
                raise
            time.sleep(.02)
            continue
        if row['state'] in {'verified', 'blocked'}:
            break
        time.sleep(.02)
    else:
        pytest.fail('Presentation operation did not reach a readable terminal state within its deadline')
    assert row['state'] == 'verified', {key: row[key] for key in ('state', 'error_code', 'result_object')}
    system[0].delta.owned_completion(admitted.job_id).result(
        timeout=max(1, deadline - time.monotonic())
    )
    return json.loads(store.lane('plan').read_object(row['result_object']))['result']['result']


def test_intake_retrieves_slide_notes_tables_and_exact_original_from_own_lane(presentation_system):
    plan(presentation_system, ['presentation_index'])
    store = presentation_system[1]
    original = (store.source_root / 'fixture.pptx').read_bytes()
    indexed = execute(presentation_system, 'presentation_index', {'filename': 'fixture.pptx'})
    assert indexed['sha256'] == digest(original) and '/ppt/objects/natural/' in indexed['natural_path'].replace('\\', '/')
    assert not indexed['fidelity']['layout_verified']
    before = {p: p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    arguments = {'snapshot_id': indexed['snapshot_id']}
    for collection, expected in [('slide', 2), ('notes', 2), ('table', 2)]:
        response = call(presentation_system, 'presentation_query', arguments | {'collection': collection})
        assert response.status == 'ok', response.error
        assert len(response.result['result']['rows']) == expected
    found = call(presentation_system, 'presentation_query', arguments | {'query': 'original source'})
    assert found.status == 'ok' and len(found.result['result']['rows']) == 2
    raw = call(presentation_system, 'presentation_read', arguments | {'max_bytes': 131072})
    assert base64.b64decode(raw.result['result']['content_base64']) == original
    assert before == {p: p.read_bytes() for p in store.root.rglob('*.sqlite*') if p.is_file()}
    with store.connection(read_only=True) as db:
        assert db.execute("SELECT count(*) FROM sqlite_schema WHERE name GLOB 'ppt_*'").fetchone()[0] == 0
    with store.lane('ppt').connection(read_only=True) as db:
        assert db.execute('SELECT count(*) FROM ppt_table').fetchone()[0] == 2
    assert (store.source_root / 'fixture.pptx').read_bytes() == original


def test_edit_reorder_and_export_preserve_original_and_unedited_members(presentation_system):
    plan(presentation_system, ['presentation_index', 'presentation_edit', 'presentation_export', 'presentation_refresh'])
    store = presentation_system[1]
    original = (store.source_root / 'fixture.pptx').read_bytes()
    indexed = execute(presentation_system, 'presentation_index', {'filename': 'fixture.pptx'})
    edited = execute(presentation_system, 'presentation_edit', {'snapshot_id': indexed['snapshot_id'], 'expected_sha256': indexed['sha256'],
        'slide_order': ['ppt/slides/slide2.xml', 'ppt/slides/slide1.xml'], 'replacements': [
            {'part': 'ppt/slides/slide1.xml', 'paragraph_index': 1,
             'expected_text': 'The selected operation preserves the original source.', 'replacement_text': 'A separate edited version.'}]}, index=1)
    changed = store.lane('ppt').read_object(edited['sha256'])
    assert edited['generation'] == 2 and edited['tool_evidence']['untouched_members_byte_identical']
    assert (store.source_root / 'fixture.pptx').read_bytes() == original
    for name, raw in package_members(original).items():
        if name not in {'ppt/slides/slide1.xml', 'ppt/presentation.xml'}:
            assert package_members(changed)[name] == raw
    from pptx import Presentation
    independently_read = Presentation(io.BytesIO(changed))
    assert independently_read.slides[0].shapes[0].text == 'Versioned presentation 2'
    assert independently_read.slides[1].shapes[1].text == 'A separate edited version.'
    assert independently_read.slides[1].notes_slide.notes_text_frame.text == 'Speaker evidence for slide 1.'
    exported = execute(presentation_system, 'presentation_export', {'snapshot_id': edited['snapshot_id'],
        'filename': 'fixture.pptx', 'expected_sha256': indexed['sha256']}, index=2)
    assert exported['source_bytes_mutated'] and (store.source_root / 'fixture.pptx').read_bytes() == changed
    destination = exported['index_refresh']['result']
    assert not exported['source_index_refresh_required'] and destination['generation'] == 3
    refreshed = execute(presentation_system, 'presentation_refresh', {'filename': 'fixture.pptx', 'expected_snapshot': destination['snapshot_id']}, index=3)
    assert refreshed['sha256'] == edited['sha256'] and refreshed['generation'] == 4
    old = call(presentation_system, 'presentation_read', {'snapshot_id': indexed['snapshot_id'], 'max_bytes': 131072})
    assert base64.b64decode(old.result['result']['content_base64']) == original


def test_generation_is_editable_in_independent_reader(presentation_system):
    plan(presentation_system, ['presentation_generate'])
    generated = execute(presentation_system, 'presentation_generate', generation(12))
    from pptx import Presentation
    raw = presentation_system[1].lane('ppt').read_object(generated['sha256'])
    independently_read = Presentation(io.BytesIO(raw))
    assert len(independently_read.slides) == 12
    assert independently_read.slides[9].shapes[0].text == 'Versioned presentation 10'
    assert independently_read.slides[0].shapes[2].table.cell(2, 1).text == 'New immutable version'
    from evidence_lane_plugin.presentation_profile import read_snapshot
    _, facts = read_snapshot(presentation_system[1], generated['snapshot_id'])
    assert facts['slide_order'][9] == 'ppt/slides/slide10.xml'
    assert not generated['fidelity']['layout_verified']
