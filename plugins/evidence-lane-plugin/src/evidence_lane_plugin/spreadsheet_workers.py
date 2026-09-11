"""Spreadsheet tool bodies called only through bounded owned OS workers."""
from __future__ import annotations

import base64
import datetime as dt
import decimal
import io
import math
import zipfile
from pathlib import PurePosixPath

from .document_parsers import digest, package_members, xml_root
from .spreadsheet_parsers import OPENXML, S, coordinates, parse_openxml
from .tabular_values import MAX_COLUMNS, MAX_FILE_BYTES, MAX_ROWS, MAX_TABLES, Facts, typed


def column_label(number):
    text = ''
    while number:
        number, remainder = divmod(number - 1, 26)
        text = chr(65 + remainder) + text
    return text


def parse_calamine(filename, content):
    from python_calamine import CalamineWorkbook, SheetTypeEnum
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in {'.xls', '.xlsb', '.ods'}:
        raise ValueError('SPREADSHEET_CALAMINE_FORMAT_UNSUPPORTED')
    if extension in {'.ods', '.xlsb'}:
        package_members(content)
    facts = Facts('data_excel', extension)
    complete = True
    with CalamineWorkbook.from_filelike(io.BytesIO(content)) as workbook:
        if len(workbook.sheet_names) > MAX_TABLES:
            raise ValueError('SPREADSHEET_SHEET_BUDGET')
        facts.add('workbook', 'workbook', 0, PurePosixPath(filename).name, sheet_count=len(workbook.sheet_names),
                  date_system='reader_converted_date_values')
        for index, metadata in enumerate(workbook.sheets_metadata):
            if metadata.typ != SheetTypeEnum.WorkSheet:
                facts.add('sheet', metadata.name, index, metadata.name, name=metadata.name, state=str(metadata.visible),
                          sheet_type=str(metadata.typ), rows_extracted=False, complete=False)
                complete = False
                continue
            sheet = workbook.get_sheet_by_index(index)
            if sheet.width > MAX_COLUMNS:
                raise ValueError('SPREADSHEET_COLUMN_BUDGET')
            start = sheet.start or (0, 0)
            rows = sheet.to_python(skip_empty_area=True, nrows=MAX_ROWS)
            selected_complete = sheet.height <= MAX_ROWS
            complete &= selected_complete
            facts.add('sheet', metadata.name, index, metadata.name, name=metadata.name, state=str(metadata.visible),
                sheet_type=str(metadata.typ), start=list(start), end=list(sheet.end) if sheet.end else None,
                sampled_rows=len(rows), complete=selected_complete)
            ordinal = 0
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    reference = column_label(start[1] + column_index + 1) + str(start[0] + row_index + 1)
                    coordinates(reference)
                    parsed = typed(value)
                    facts.add('cell', metadata.name, ordinal, str(parsed['value']) if parsed['value'] is not None else '',
                        sheet=metadata.name, cell=reference, row=start[0] + row_index + 1, column=start[1] + column_index + 1,
                        value=parsed, formula=None, formula_attributes={}, cached_value=None, style=None,
                        expression_availability='not_extracted_by_value_reader')
                    ordinal += 1
    return facts.finish(fidelity={'structure': 'calamine_sheet_metadata_and_values', 'complete': complete,
        'formulas': 'expressions_not_extracted_values_may_be_cached', 'layout': 'not_rendered'}, limitations=[
        'XLS/XLSB/ODS intake retains original bytes and bounded reader-converted cell values.',
        'Formula expressions, charts and detailed formatting are not extracted by this value reader.',
        'Blank cells may be reader-normalized empty strings; this is not native OpenXML cell typing.',
        'No macros, external links or formulas are executed.'])


def parse_spreadsheet(filename, content):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('TABULAR_FILE_BYTE_BUDGET')
    return parse_openxml(filename, content) if PurePosixPath(filename).suffix.lower() in OPENXML else parse_calamine(filename, content)


def _python_value(value):
    kind, body = value['type'], value['value']
    if kind in {'text', 'boolean', 'null'}:
        return body
    if kind == 'number':
        return _excel_number(body)
    if kind in {'date', 'iso_date'}:
        return dt.date.fromisoformat(body)
    if kind == 'datetime':
        return dt.datetime.fromisoformat(body)
    if kind == 'time':
        return dt.time.fromisoformat(body)
    raise ValueError('SPREADSHEET_GENERATE_VALUE_TYPE_UNSUPPORTED')


def _excel_number(body):
    value = decimal.Decimal(body)
    digits = ''.join(str(digit) for digit in value.as_tuple().digits).rstrip('0')
    if (not value.is_finite() or len(digits) > 15 or not math.isfinite(float(value))
            or value != 0 and not decimal.Decimal('2.22507385850721e-308') <= abs(value) <= decimal.Decimal('9.99999999999999e307')):
        raise ValueError('SPREADSHEET_NUMBER_FIDELITY_UNSUPPORTED_USE_TEXT')
    return value


def _range_contains(reference, cell):
    target_row, target_column = coordinates(cell)
    for selected in reference.split():
        ends = selected.replace('$', '').split(':')
        if len(ends) not in {1, 2}:
            raise ValueError('SPREADSHEET_EDIT_RANGE_UNSUPPORTED')
        first, last = coordinates(ends[0]), coordinates(ends[-1])
        if first[0] > last[0] or first[1] > last[1]:
            raise ValueError('SPREADSHEET_EDIT_RANGE_UNSUPPORTED')
        if first[0] <= target_row <= last[0] and first[1] <= target_column <= last[1]:
            return True
    return False


def _editable_cell(root, node, reference, members):
    protection = root.find(S + 'sheetProtection')
    if protection is not None and protection.get('sheet', '0') in {'1', 'true'}:
        raise ValueError('SPREADSHEET_EDIT_PROTECTED_SHEET')
    string = node.find(S + 'is')
    if node.get('t') == 's':
        string = xml_root(members['xl/sharedStrings.xml']).findall(S + 'si')[int(node.find(S + 'v').text)]
    if string is not None and any(child.tag != S + 't' for child in string):
        raise ValueError('SPREADSHEET_EDIT_RICH_TEXT_UNSUPPORTED')
    if root.find(S + 'extLst') is not None or any(key in node.attrib for key in ('cm', 'vm')):
        raise ValueError('SPREADSHEET_EDIT_EXTENDED_CELL_METADATA_UNSUPPORTED')
    for formula in root.iter(S + 'f'):
        if formula.get('ref') and _range_contains(formula.get('ref'), reference):
            raise ValueError('SPREADSHEET_EDIT_COMPLEX_FORMULA_RANGE_UNSUPPORTED')
    for tag, attribute in (('mergeCell', 'ref'), ('dataValidation', 'sqref')):
        for selected in root.iter(S + tag):
            if _range_contains(selected.get(attribute, ''), reference):
                raise ValueError('SPREADSHEET_EDIT_MERGED_OR_VALIDATED_RANGE_UNSUPPORTED')


def generate(arguments):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    workbook = Workbook(iso_dates=True)
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    for selected in arguments['sheets']:
        worksheet = workbook.create_sheet(selected['name'])
        seen = set()
        for entry in selected['cells']:
            row, column = coordinates(entry['cell'])
            if entry['cell'] in seen:
                raise ValueError('SPREADSHEET_GENERATE_DUPLICATE_CELL')
            seen.add(entry['cell'])
            cell = worksheet.cell(row, column)
            if entry.get('formula'):
                if entry['value']['type'] != 'null':
                    raise ValueError('SPREADSHEET_GENERATE_FORMULA_CACHE_NOT_ACCEPTED')
                cell.value = '=' + entry['formula'].lstrip('=')
            else:
                selected_value = _python_value(entry['value'])
                if isinstance(selected_value, (dt.time, dt.datetime)) and (selected_value.tzinfo is not None or selected_value.microsecond % 1000):
                    raise ValueError('SPREADSHEET_DATE_TIME_FIDELITY_UNSUPPORTED')
                cell.value = selected_value
                if entry['value']['type'] == 'text':
                    cell.data_type = 's'
            if entry.get('number_format'):
                cell.number_format = entry['number_format']
            cell.font = Font(name='Calibri', size=11, color='172C3D')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row == 1:
                cell.font = Font(name='Calibri', size=11, color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='173B56')
            worksheet.column_dimensions[column_label(column)].width = 22
        worksheet.freeze_panes = selected.get('freeze_panes')
        if selected.get('autofilter'):
            worksheet.auto_filter.ref = selected['autofilter']
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_view.showGridLines = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()
    parsed = parse_openxml(arguments['logical_name'], content)
    cells = {(item['sheet'], item['cell']): item for item in parsed['items'] if item['kind'] == 'cell'}
    for sheet in arguments['sheets']:
        for entry in sheet['cells']:
            actual = cells.get((sheet['name'], entry['cell']))
            if actual is None:
                raise ValueError('SPREADSHEET_GENERATE_CELL_ROUNDTRIP_FAILED')
            value = entry['value']
            if value['type'] == 'number' and (actual['value']['type'] != 'number' or decimal.Decimal(actual['value']['value']) != decimal.Decimal(value['value'])):
                raise ValueError('SPREADSHEET_GENERATE_NUMBER_ROUNDTRIP_FAILED')
            if value['type'] in {'text', 'boolean'} and actual['value'] != value:
                raise ValueError('SPREADSHEET_GENERATE_VALUE_ROUNDTRIP_FAILED')
    from .tabular_workers import encoded
    return encoded('data_excel', arguments['logical_name'], content, evidence={'tool': 'openpyxl',
        'formulas_recalculated': False, 'cached_values': 'not_generated', 'layout_verified': False})


def _set_cell(node, value, formula, ET):
    for child in list(node):
        if child.tag in {S + 'v', S + 'f', S + 'is'}:
            node.remove(child)
    node.attrib.pop('t', None)
    if formula:
        if value['type'] != 'null':
            raise ValueError('SPREADSHEET_EDIT_FORMULA_CACHE_NOT_ACCEPTED')
        ET.SubElement(node, S + 'f').text = formula.lstrip('=')
    elif value['type'] == 'text':
        node.set('t', 'inlineStr')
        text = ET.SubElement(ET.SubElement(node, S + 'is'), S + 't')
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        text.text = value['value']
    elif value['type'] == 'number':
        ET.SubElement(node, S + 'v').text = str(_excel_number(value['value']))
    elif value['type'] == 'boolean':
        node.set('t', 'b')
        ET.SubElement(node, S + 'v').text = '1' if value['value'] else '0'
    elif value['type'] == 'null':
        pass
    else:
        raise ValueError('SPREADSHEET_EDIT_VALUE_TYPE_UNSUPPORTED')


def edit(arguments):
    from lxml import etree as ET
    content = base64.b64decode(arguments['content_base64'], validate=True)
    if digest(content) != arguments['expected_sha256']:
        raise ValueError('SPREADSHEET_EDIT_HASH_MISMATCH')
    if PurePosixPath(arguments['logical_name']).suffix.lower() not in OPENXML:
        raise ValueError('SPREADSHEET_EDIT_FORMAT_UNSUPPORTED')
    before = parse_openxml(arguments['logical_name'], content)
    members = package_members(content)
    if any(name.startswith('_xmlsignatures/') for name in members):
        raise ValueError('SPREADSHEET_EDIT_SIGNED_PACKAGE_UNSUPPORTED')
    protection = xml_root(members['xl/workbook.xml']).find(S + 'workbookProtection')
    if protection is not None and any(protection.get(flag, '0') in {'1', 'true'} for flag in ('lockStructure', 'lockWindows', 'lockRevision')):
        raise ValueError('SPREADSHEET_EDIT_PROTECTED_WORKBOOK')
    lookup = {(item['sheet'], item['cell']): item for item in before['items'] if item['kind'] == 'cell'}
    sheets = {item['name']: item['part'] for item in before['items'] if item['kind'] == 'sheet'}
    roots, changes, selected = {}, [], set()
    for change in arguments['replacements']:
        identity = change['sheet'], change['cell']
        if identity in selected or change['sheet'] not in sheets:
            raise ValueError('SPREADSHEET_EDIT_LOCATOR_INVALID')
        selected.add(identity)
        item = lookup.get(identity)
        if item is None or item['value'] != change['expected_value'] or item['formula'] != change.get('expected_formula'):
            raise ValueError('SPREADSHEET_EDIT_CELL_CHANGED_OR_ABSENT')
        if item['formula_attributes']:
            raise ValueError('SPREADSHEET_EDIT_COMPLEX_FORMULA_UNSUPPORTED')
        part = item['part']
        if part not in roots:
            xml_root(members[part])
            roots[part] = ET.fromstring(members[part], parser=ET.XMLParser(resolve_entities=False, no_network=True))
        node = next(node for node in roots[part].iter(S + 'c') if node.get('r') == change['cell'])
        _editable_cell(roots[part], node, change['cell'], members)
        if any(child.tag not in {S + 'v', S + 'f', S + 'is'} for child in node):
            raise ValueError('SPREADSHEET_EDIT_COMPLEX_CELL_UNSUPPORTED')
        _set_cell(node, change['replacement_value'], change.get('replacement_formula'), ET)
        changes.append({'sheet': change['sheet'], 'cell': change['cell'], 'style_preserved': True})
    cleared = 0
    # Any changed input can invalidate formula caches on other sheets. Preserve
    # expressions and styles while making this uncertainty explicit in the file.
    for part in sheets.values():
        if part not in roots:
            roots[part] = ET.fromstring(members[part], parser=ET.XMLParser(resolve_entities=False, no_network=True))
        changed = part in {item['part'] for key, item in lookup.items() if key in selected}
        for cell in roots[part].iter(S + 'c'):
            if cell.find(S + 'f') is not None:
                for cached in list(cell.findall(S + 'v')):
                    cell.remove(cached)
                    cleared += 1
                    changed = True
        if not changed:
            roots.pop(part)
    workbook = ET.fromstring(members['xl/workbook.xml'], parser=ET.XMLParser(resolve_entities=False, no_network=True))
    calc = workbook.find(S + 'calcPr')
    if calc is None:
        calc = ET.SubElement(workbook, S + 'calcPr')
    calc.attrib.update({'fullCalcOnLoad': '1', 'forceFullCalc': '1', 'calcMode': 'auto'})
    roots['xl/workbook.xml'] = workbook
    for part, root in roots.items():
        members[part] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(output, 'w') as edited:
        for info in original.infolist():
            edited.writestr(info, members.get(info.filename, b''))
    raw = output.getvalue()
    after_members = package_members(raw)
    if any(raw != after_members[name] for name, raw in package_members(content).items() if name not in roots):
        raise ValueError('SPREADSHEET_EDIT_UNTOUCHED_MEMBER_CHANGED')
    from .tabular_workers import encoded
    return encoded('data_excel', arguments['logical_name'], raw, evidence={'tool': 'native_OpenXML_lxml', 'changes': changes,
        'changed_parts': sorted(roots), 'untouched_members_byte_identical': True,
        'formula_caches_cleared': cleared, 'formulas_recalculated': False, 'layout_verified': False})
