"""Executable Excel, DataFrame, SQLAlchemy, and Tableau Hyper adapters."""

from __future__ import annotations

import importlib.util
from pathlib import Path
from typing import Any

from pydantic import BaseModel, ConfigDict, Field

from .hashing import canonical_json_bytes, sha256_bytes, sha256_file


class DataInspectionRequest(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    source_path: Path
    host_profile: str
    max_rows: int = Field(default=200, ge=1, le=2_000)


def _codex_host(value: str) -> str:
    exact = value.strip().upper()
    if exact not in {"CODEX_DESKTOP", "CODEX_CLI", "CODEX_VM"}:
        raise ValueError("DATA_TOOLCHAIN_CODEX_HOST_PROFILE_REQUIRED")
    return exact


def inspect_excel_openpyxl(request: DataInspectionRequest) -> dict[str, Any]:
    path = request.source_path.resolve(strict=True)
    host = _codex_host(request.host_profile)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("OPENPYXL_WORKBOOK_FORMAT_REQUIRED")
    if importlib.util.find_spec("openpyxl") is None:
        raise RuntimeError("OPENPYXL_DEPENDENCY_UNAVAILABLE")
    import openpyxl  # type: ignore[import-not-found]

    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=True,
    )
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            formulas = 0
            nonempty = 0
            samples: list[dict[str, Any]] = []
            for row_index, row in enumerate(worksheet.iter_rows(), start=1):
                if row_index > request.max_rows:
                    break
                for cell in row:
                    if cell.value is None:
                        continue
                    nonempty += 1
                    if cell.data_type == "f":
                        formulas += 1
                    if len(samples) < 500:
                        samples.append(
                            {
                                "coordinate": cell.coordinate,
                                "value": str(cell.value),
                                "data_type": str(cell.data_type),
                            }
                        )
            sheets.append(
                {
                    "title": worksheet.title,
                    "max_row": int(worksheet.max_row or 0),
                    "max_column": int(worksheet.max_column or 0),
                    "bounded_nonempty_cells": nonempty,
                    "bounded_formula_cells": formulas,
                    "samples": samples,
                }
            )
    finally:
        workbook.close()
    core = {
        "schema": "evidence-lane.openpyxl-inspection.v1",
        "status": "PASS",
        "engine": "openpyxl",
        "engine_version": str(openpyxl.__version__),
        "host_profile": host,
        "source_sha256": sha256_file(path),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "read_only": True,
        "workbook_mutated": False,
    }
    return {**core, "receipt_sha256": sha256_bytes(canonical_json_bytes(core))}


def inspect_tabular_pandas(request: DataInspectionRequest) -> dict[str, Any]:
    path = request.source_path.resolve(strict=True)
    host = _codex_host(request.host_profile)
    if importlib.util.find_spec("pandas") is None:
        raise RuntimeError("PANDAS_DEPENDENCY_UNAVAILABLE")
    import pandas as pd  # type: ignore[import-not-found]

    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        sheets = pd.read_excel(path, sheet_name=None, nrows=request.max_rows)
    elif suffix in {".csv", ".tsv"}:
        frame = pd.read_csv(
            path,
            sep="\t" if suffix == ".tsv" else ",",
            nrows=request.max_rows,
        )
        sheets = {"table": frame}
    elif suffix == ".parquet":
        frame = pd.read_parquet(path).head(request.max_rows)
        sheets = {"table": frame}
    else:
        raise ValueError("PANDAS_SOURCE_FORMAT_UNSUPPORTED")
    projections = []
    for name, frame in sheets.items():
        projections.append(
            {
                "name": str(name),
                "bounded_rows": len(frame.index),
                "columns": [str(value) for value in frame.columns],
                "dtypes": {str(key): str(value) for key, value in frame.dtypes.items()},
                "null_counts": {
                    str(key): int(value) for key, value in frame.isna().sum().items()
                },
            }
        )
    core = {
        "schema": "evidence-lane.pandas-inspection.v1",
        "status": "PASS",
        "engine": "pandas",
        "engine_version": str(pd.__version__),
        "host_profile": host,
        "source_sha256": sha256_file(path),
        "projections": projections,
        "bounded": True,
        "source_mutated": False,
    }
    return {**core, "receipt_sha256": sha256_bytes(canonical_json_bytes(core))}


def inspect_sqlalchemy_sqlite(
    request: DataInspectionRequest,
) -> dict[str, Any]:
    path = request.source_path.resolve(strict=True)
    host = _codex_host(request.host_profile)
    if path.suffix.lower() not in {".sqlite", ".sqlite3", ".db"}:
        raise ValueError("SQLALCHEMY_SQLITE_SOURCE_REQUIRED")
    if importlib.util.find_spec("sqlalchemy") is None:
        raise RuntimeError("SQLALCHEMY_DEPENDENCY_UNAVAILABLE")
    from sqlalchemy import create_engine, inspect  # type: ignore[import-not-found]

    engine = create_engine(f"sqlite+pysqlite:///file:{path.as_posix()}?mode=ro&uri=true")
    try:
        inspector = inspect(engine)
        tables = []
        for name in sorted(inspector.get_table_names()):
            tables.append(
                {
                    "name": name,
                    "columns": [
                        {
                            "name": str(row["name"]),
                            "type": str(row["type"]),
                            "nullable": bool(row.get("nullable", True)),
                        }
                        for row in inspector.get_columns(name)
                    ],
                    "foreign_keys": inspector.get_foreign_keys(name),
                    "indexes": inspector.get_indexes(name),
                }
            )
        views = sorted(inspector.get_view_names())
    finally:
        engine.dispose()
    core = {
        "schema": "evidence-lane.sqlalchemy-inspection.v1",
        "status": "PASS",
        "engine": "SQLAlchemy",
        "host_profile": host,
        "source_sha256": sha256_file(path),
        "tables": tables,
        "views": views,
        "read_only_uri": True,
        "source_mutated": False,
    }
    return {**core, "receipt_sha256": sha256_bytes(canonical_json_bytes(core))}


def inspect_tableau_hyper(request: DataInspectionRequest) -> dict[str, Any]:
    path = request.source_path.resolve(strict=True)
    host = _codex_host(request.host_profile)
    if path.suffix.lower() != ".hyper":
        raise ValueError("TABLEAU_HYPER_SOURCE_REQUIRED")
    if importlib.util.find_spec("tableauhyperapi") is None:
        raise RuntimeError("TABLEAU_HYPER_API_DEPENDENCY_UNAVAILABLE")
    from tableauhyperapi import (  # type: ignore[import-not-found]
        Connection,
        CreateMode,
        HyperProcess,
        Telemetry,
    )

    tables: list[dict[str, Any]] = []
    with (
        HyperProcess(Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU) as process,
        Connection(
            process.endpoint,
            str(path),
            create_mode=CreateMode.NONE,
        ) as connection,
    ):
        for table in connection.catalog.get_table_names("Extract"):
            definition = connection.catalog.get_table_definition(table)
            count = int(connection.execute_scalar_query(f"SELECT COUNT(*) FROM {table}"))
            samples = [
                [str(value) if value is not None else None for value in row]
                for row in connection.execute_list_query(
                    f"SELECT * FROM {table} LIMIT {int(request.max_rows)}"
                )
            ]
            tables.append(
                {
                    "name": str(table),
                    "columns": [
                        {"name": str(column.name), "type": str(column.type)}
                        for column in definition.columns
                    ],
                    "row_count": count,
                    "bounded_rows": samples,
                }
            )
    core = {
        "schema": "evidence-lane.tableau-hyper-inspection.v1",
        "status": "PASS",
        "engine": "tableauhyperapi",
        "host_profile": host,
        "source_sha256": sha256_file(path),
        "tables": tables,
        "telemetry_enabled": False,
        "source_mutated": False,
    }
    return {**core, "receipt_sha256": sha256_bytes(canonical_json_bytes(core))}


__all__ = [
    "DataInspectionRequest",
    "inspect_excel_openpyxl",
    "inspect_sqlalchemy_sqlite",
    "inspect_tableau_hyper",
    "inspect_tabular_pandas",
]
